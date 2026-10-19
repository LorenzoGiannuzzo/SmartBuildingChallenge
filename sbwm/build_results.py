"""Build the ``results`` folder: figures (PDF, JPG) and tables (CSV, Excel).

Layout:
    results/
      pdf/   jpg/   csv/   excel/
        <general figures and tables>
        building_<n>/   <figures and tables of one building>

Steps:
  1. Simulate every shipped scenario with the DPC policy and with PID With
     Water Reset for ``--days`` days from the official start date (same
     weather and seeded occupancy for both), saving 5-minute trajectories.
  2. Pre-conditioning check: energy spent before occupants arrive, on
     non-working days, and how early the plant starts before comfort hours.
  3. Tables and figures, general and per building.

Usage (from the repository root):
    python -m sbwm.build_results --weights policies/dpc_v3.pt --days 7 --workers 8
    python -m sbwm.build_results --skip-sim      # tables and figures only
"""

from __future__ import annotations

import argparse
import csv
import json
import multiprocessing as mp
import re
import shutil
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from sbwm import make_figures as mf  # noqa: E402

ROOT = Path("results")
SCEN_ROOT = Path("smart_buildings/training_scenarios")
CONTROLLERS = {"dpc": "DPC Transformer (Ours)", "pid_water": "PID With Water Reset"}
COLORS = {"dpc": mf.BLUE, "pid_water": mf.ORANGE}
PRE_HOURS = 3  # pre-arrival window used by the check [h]


def building_dir(building_key: str) -> str:
    return "building_" + building_key.split("_")[-1]


def climate_name(weather_key: str) -> str:
    return mf_climate(weather_key.split("_")[0])


def mf_climate(station: str) -> str:
    return station.capitalize()


def ensure_dirs(buildings):
    for fmt in ("pdf", "jpg", "csv", "excel"):
        (ROOT / fmt).mkdir(parents=True, exist_ok=True)
        for b in buildings:
            (ROOT / fmt / b).mkdir(parents=True, exist_ok=True)


def make_saver(sub: str = ""):
    def saver(fig, name):
        for fmt, kw in (("pdf", {}), ("jpg", {"pil_kwargs": {"quality": 92}})):
            path = ROOT / fmt / sub / f"{name}.{fmt}"
            path.parent.mkdir(parents=True, exist_ok=True)
            fig.savefig(path, **kw)
        print("wrote", ROOT / "pdf" / sub / f"{name}.pdf", flush=True)
    return saver


# ---------------------------------------------------------------- simulation

def _sim_job(job):
    try:
        import torch

        from sbwm.record_timeseries import record

        torch.set_num_threads(1)
        rows = record(Path(job["scenario"]), job["controller"], job["weights"], job["steps"])
        out = Path(job["out"])
        with open(out, "w", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        return f"done {out}"
    except Exception:  # noqa: BLE001
        import traceback

        return f"FAILED {job['out']}\n{traceback.format_exc()}"


def simulate(scenarios, weights, steps, workers):
    jobs = []
    for sc in scenarios:
        b, w = sc.parent.name, sc.name
        for ctrl in CONTROLLERS:
            out = ROOT / "csv" / building_dir(b) / f"timeseries_{w}_{ctrl}.csv"
            #Lorenzo Giannuzzo: an existing trajectory is reused only if it has
            #the requested length, so changing --days re-simulates.
            if out.exists() and sum(1 for _ in open(out)) - 1 == steps:
                continue
            jobs.append({"scenario": str(sc), "controller": ctrl, "weights": weights,
                         "steps": steps, "out": str(out)})
    print(f"{len(jobs)} simulations to run", flush=True)
    if not jobs:
        return
    with mp.get_context("spawn").Pool(workers, maxtasksperchild=1) as pool:
        for msg in pool.imap_unordered(_sim_job, jobs):
            print(msg, flush=True)


# ------------------------------------------------------- pre-conditioning check

def precondition_metrics(df: pd.DataFrame) -> dict:
    """Energy split around comfort hours, from one 5-minute trajectory."""
    t = pd.to_datetime(df["time"])
    power = df["electricity_kw"] + df["gas_kw"]  # kW
    kwh = power * 5 / 60
    eco = df["heating_setpoint"].min()
    comfort = df["heating_setpoint"] > eco + 0.5
    day = t.dt.date
    working = comfort.groupby(day).transform("any")
    pre = pd.Series(False, index=df.index)
    leads = []
    for d, idx in df.groupby(day).groups.items():
        idx = np.asarray(idx)
        c = comfort.values[idx]
        if not c.any():
            continue
        first = idx[np.argmax(c)]
        window = (t >= t[first] - pd.Timedelta(hours=PRE_HOURS)) & (t < t[first])
        pre |= window
        #Lorenzo Giannuzzo: lead time is how long before comfort start the plant
        #is continuously running above 10% of its mean comfort-hour power.
        thr = 0.1 * power[comfort].mean() if comfort.any() else 0.0
        k, lead = first - 1, 0.0
        while k >= 0 and power[k] > thr and lead < 12:
            lead += 5 / 60
            k -= 1
        leads.append(lead)
    total = kwh.sum()
    occupied = kwh[comfort].sum()
    pre_e = kwh[pre & ~comfort].sum()
    nonwork = kwh[~working].sum()
    other = total - occupied - pre_e - nonwork
    share = (lambda x: 100 * x / total) if total > 0 else (lambda x: 0.0)
    return {
        "Total Energy [kWh]": round(total, 1),
        "Comfort Hours Energy [kWh]": round(occupied, 1),
        f"Pre-Arrival Energy, {PRE_HOURS} H Before Comfort Start [kWh]": round(pre_e, 1),
        "Other Unoccupied Hours Energy [kWh]": round(other, 1),
        "Non-Working Days Energy [kWh]": round(nonwork, 1),
        "Pre-Arrival Share [%]": round(share(pre_e), 1),
        "Non-Working Days Share [%]": round(share(nonwork), 1),
        "Median Start Lead Before Comfort Hours [h]": round(float(np.median(leads)), 2) if leads else 0.0,
        "Non-Working Days (Count)": int((~working.groupby(day).first()).sum()),
        "Mean Reward Per Step [-]": round(df["reward"].mean(), 5),
    }


# --------------------------------------------------------------------- tables

def official_tables(out_csv: Path):
    """Official 30-day results and baselines, via sbwm.report_tables."""
    from sbwm import report_tables as rt

    rt.official_table(Path("evaluations/official_eval_dpc_v3.txt"), out_csv)
    for name in ("loop_benchmark_6_scenarios_2_days.csv", "world_model_accuracy.csv"):
        src = Path("evaluations") / name
        if src.exists():
            shutil.copy(src, out_csv / name)


def write_excel(path: Path, sheets: dict[str, pd.DataFrame]):
    try:
        import openpyxl  # noqa: F401
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed, Excel files skipped (pip install openpyxl)")
        return
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for j, col in enumerate(df.columns, start=1):
                width = max(len(str(col)), *(len(str(v)) for v in df[col].head(200))) + 2
                ws.column_dimensions[get_column_letter(j)].width = min(width, 45)
                #Lorenzo Giannuzzo: plain number formats, no thousands separator.
                if pd.api.types.is_float_dtype(df[col]):
                    for cell in ws[get_column_letter(j)][1:]:
                        cell.number_format = "0.0000" if df[col].abs().max() < 1 else "0.0"
    print("wrote", path, flush=True)


# -------------------------------------------------------------------- figures

def hourly(df):
    g = df.groupby(np.arange(len(df)) // 12)
    return pd.DataFrame({
        "temp_mean": g["temp_mean"].mean(), "temp_min": g["temp_min"].mean(),
        "temp_max": g["temp_max"].mean(), "heating_setpoint": g["heating_setpoint"].min(),
        "cooling_setpoint": g["cooling_setpoint"].max(), "occupants": g["occupants"].mean(),
        "power": (g["electricity_kw"].mean() + g["gas_kw"].mean())})


def scenario_figure(runs: dict[str, pd.DataFrame], title: str, name: str, sub: str):
    hd = {k: hourly(v) for k, v in runs.items()}
    ref = hd["dpc"]
    n = len(ref)
    fig, axes = plt.subplots(3, 1, figsize=(7.6, 8.4), sharex=True,
                             gridspec_kw={"height_ratios": [3, 2, 2], "hspace": 0.12})
    ax1, ax2, ax3 = axes
    occupied = ref["occupants"].values > 0.5
    edges = np.flatnonzero(np.diff(np.r_[0, occupied.astype(int), 0]))
    for ax in axes:
        for s0, s1 in zip(edges[::2], edges[1::2]):
            ax.axvspan(s0, s1, color="#f7f3e6", zorder=0, lw=0)
    x_edges = np.arange(n + 1)
    ax1.fill_between(x_edges, np.r_[ref["heating_setpoint"], ref["heating_setpoint"].iloc[-1]],
                     np.r_[ref["cooling_setpoint"], ref["cooling_setpoint"].iloc[-1]], step="post",
                     color=mf.BAND, edgecolor=mf.MUTED, lw=1, zorder=1, label="Comfort Band")
    h = np.arange(n) + 0.5
    for k in ("pid_water", "dpc"):
        d = hd[k]
        ax1.fill_between(h, d["temp_min"], d["temp_max"], color=COLORS[k], alpha=0.12, lw=0, zorder=2)
        ax1.plot(h, d["temp_mean"], color=COLORS[k], lw=2, zorder=3, label=CONTROLLERS[k])
        ax2.plot(h, d["power"], color=COLORS[k], lw=2, label=CONTROLLERS[k])
        raw = runs[k]
        ax3.plot(np.arange(len(raw)) * 5 / 60, raw["reward"].cumsum(), color=COLORS[k], lw=2)
        ax3.text(len(raw) * 5 / 60 + 1, raw["reward"].sum(), f"{raw['reward'].sum():.1f}",
                 va="center", fontsize=9, color=mf.INK)
    lo = min(min(d["temp_min"].min() for d in hd.values()), ref["heating_setpoint"].min())
    hi = max(max(d["temp_max"].max() for d in hd.values()), ref["cooling_setpoint"].max())
    ax1.set_ylim(lo - 1, hi + (hi - lo) * 0.45)
    ax1.set_yticks([t for t in ax1.get_yticks() if lo - 1 <= t <= hi + 0.5])
    ax1.set_ylabel("Zone Temperature [°C]")
    mf._legend(ax1, 3)
    ax1.set_title(title, loc="left", color=mf.INK, fontsize=11)
    ax2.set_ylabel("HVAC Power [kW]")
    top = max(d["power"].max() for d in hd.values())
    ax2.set_ylim(0, top * 1.1 if top > 0 else 1)
    ax3.set_ylabel("Cumulative Reward [-]")
    ax3.set_xlabel("Time [h]")
    ax3.set_xlim(0, n + max(3, 0.07 * n))
    fig.text(0.125, 0.012, "Shaded Columns: Occupied Hours. Temperature Lines: Hourly Building Mean; "
             "Bands: Zone Range. Power: Electricity Plus Gas.", fontsize=8, color=mf.INK2)
    mf.SAVER = make_saver(sub)
    mf._save(fig, name)


def building_figure(table: pd.DataFrame, building: str, sub: str):
    t = table.sort_values("Climate")
    fig, ax = plt.subplots(figsize=(7.0, 3.6))
    y = np.arange(len(t))
    hgt = 0.36
    ax.barh(y - hgt / 2 - 0.02, t["PID With Water Reset"], height=hgt, color=mf.ORANGE,
            label="PID With Water Reset")
    ax.barh(y + hgt / 2 + 0.02, t["DPC Transformer"], height=hgt, color=mf.BLUE,
            label="DPC Transformer (Ours)")
    for yi, a, b in zip(y, t["DPC Transformer"], t["PID With Water Reset"]):
        ax.text(a, yi + hgt / 2 + 0.02, f" {a:.4f}", va="center", fontsize=8.5, color=mf.INK)
        ax.text(b, yi - hgt / 2 - 0.02, f" {b:.4f}", va="center", fontsize=8.5, color=mf.INK)
    ax.set_yticks(y, t["Climate"])
    ax.invert_yaxis()
    ax.set_xlabel("Score, Mean Absolute Reward Per Step [-]")
    ax.set_xlim(0, t[["DPC Transformer", "PID With Water Reset"]].values.max() * 1.25)
    ax.set_ylim(len(t) - 0.4, -1.6)
    ax.grid(axis="y", visible=False)
    ax.set_title(f"{building}: Official Score Per Climate (30 Days Each)", loc="left", color=mf.INK, fontsize=11)
    mf._legend(ax, 2)
    mf.SAVER = make_saver(sub)
    mf._save(fig, "score_per_climate")


def precondition_figure(profiles: dict[str, np.ndarray], shares: pd.DataFrame):
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(9.6, 3.8), gridspec_kw={"width_ratios": [3, 2], "wspace": 0.3})
    n = len(next(iter(profiles.values())))
    hrs = (np.arange(n) - 96) * 5 / 60  # profiles start 8 h before comfort start
    for k in ("pid_water", "dpc"):
        ax1.plot(hrs, profiles[k], color=COLORS[k], lw=2, label=CONTROLLERS[k])
    ax1.axvline(0, color=mf.INK2, lw=1)
    ax1.axvspan(-PRE_HOURS, 0, color="#f7f3e6", zorder=0, lw=0)
    ax1.set_xlabel("Time Relative To Comfort Start [h]")
    ax1.set_ylabel("Mean HVAC Power [kW]")
    top = max(p.max() for p in profiles.values())
    ax1.set_ylim(0, top * 1.35)
    ax1.set_yticks([t for t in ax1.get_yticks() if 0 <= t <= top * 1.05])
    ax1.set_title("Power Around Comfort Start, Working Days", loc="left", color=mf.INK, fontsize=10.5)
    mf._legend(ax1, 2)
    cats = ["Pre-Arrival Share [%]", "Non-Working Days Share [%]"]
    labels = [f"Pre-Arrival ({PRE_HOURS} H)", "Non-Working Days"]
    x = np.arange(len(cats))
    w = 0.34
    plotted = []
    for i, k in enumerate(("pid_water", "dpc")):
        vals = [shares.loc[shares["Controller"] == CONTROLLERS[k], c].mean() for c in cats]
        plotted += vals
        ax2.bar(x + (i - 0.5) * (w + 0.02), vals, width=w, color=COLORS[k], label=CONTROLLERS[k])
        for xi, v in zip(x, vals):
            ax2.text(xi + (i - 0.5) * (w + 0.02), v, f"{v:.1f}", ha="center", va="bottom", fontsize=8.5,
                     color=mf.INK)
    ax2.set_xticks(x, labels)
    ax2.set_ylabel("Mean Share Of Total Energy [%]")
    ax2.set_ylim(0, max(max(plotted), 1) * 1.3)
    ax2.grid(axis="x", visible=False)
    ax2.set_title("Energy Outside Occupied Hours", loc="left", color=mf.INK, fontsize=10.5)
    mf.SAVER = make_saver("")
    mf._save(fig, "fig6_preconditioning_check")


def power_profile(df: pd.DataFrame) -> np.ndarray | None:
    eco = df["heating_setpoint"].min()
    comfort = (df["heating_setpoint"] > eco + 0.5).values
    power = (df["electricity_kw"] + df["gas_kw"]).values
    starts = np.flatnonzero(np.diff(np.r_[0, comfort.astype(int)]) == 1)
    segs = [power[s - 96:s + 36] for s in starts if s >= 96 and s + 36 <= len(power)]
    return np.mean(segs, 0) if segs else None


# ----------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--weights", default="policies/dpc_v3.pt")
    ap.add_argument("--days", type=int, default=7)
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--skip-sim", action="store_true")
    ap.add_argument("--filter", nargs="*", help="substrings a scenario path must contain")
    args = ap.parse_args()
    scenarios = sorted(p for p in SCEN_ROOT.glob("*/*") if (p / "config.gin").exists())
    if args.filter:
        scenarios = [p for p in scenarios if any(f in str(p) for f in args.filter)]
    buildings = sorted({building_dir(p.parent.name) for p in scenarios}, key=lambda b: int(b.split("_")[1]))
    ensure_dirs(buildings)

    if not args.skip_sim:
        simulate(scenarios, args.weights, args.days * 288, args.workers)

    # general official tables and figures
    official_tables(ROOT / "csv")
    mf.EVAL = ROOT / "csv"
    mf.SAVER = make_saver("")
    mf.official_scores()
    mf.improvement_heatmap()
    mf.per_scenario()
    mf.loop_progress()
    mf.world_model_accuracy()

    official = pd.read_csv(ROOT / "csv" / "official_per_scenario.csv")
    pre_rows, profiles = [], {k: [] for k in CONTROLLERS}
    for sc in scenarios:
        b, w = sc.parent.name, sc.name
        sub = building_dir(b)
        runs = {}
        for k in CONTROLLERS:
            f = ROOT / "csv" / sub / f"timeseries_{w}_{k}.csv"
            if f.exists():
                runs[k] = pd.read_csv(f)
        if len(runs) < 2:
            continue
        for k, df in runs.items():
            m = precondition_metrics(df)
            pre_rows.append({"Scenario": f"{b}/{w}", "Building": "Building " + b.split("_")[-1],
                             "Climate": climate_name(w), "Controller": CONTROLLERS[k], **m})
            prof = power_profile(df)
            if prof is not None:
                profiles[k].append(prof)
        days = f"{args.days} Day" + ("s" if args.days != 1 else "")
        title = f"Building {b.split('_')[-1]}, {climate_name(w)}: {days} Under Both Controllers"
        scenario_figure(runs, title, f"timeseries_{w}", sub)

    pre = pd.DataFrame(pre_rows)
    if not pre.empty:
        pre.to_csv(ROOT / "csv" / "preconditioning_check.csv", index=False)
        if all(profiles[k] for k in CONTROLLERS):
            precondition_figure({k: np.mean(v, 0) for k, v in profiles.items()}, pre)

    # per-building tables and figures
    for b in buildings:
        bname = "Building " + b.split("_")[1]
        tb = official[official["Building"] == bname].copy()
        if tb.empty:
            continue
        building_figure(tb, bname, b)
        tb.to_csv(ROOT / "csv" / b / "official_scores.csv", index=False)
        sheets = {"Official Scores": tb}
        if not pre.empty:
            pb = pre[pre["Building"] == bname]
            pb.to_csv(ROOT / "csv" / b / "preconditioning_check.csv", index=False)
            sheets["Pre-Conditioning Check"] = pb
        for f in sorted((ROOT / "csv" / b).glob("timeseries_*_*.csv")):
            parts = f.stem.split("_")
            sheet = f"{parts[1].capitalize()} {parts[2]} " + ("DPC" if parts[-1] == "dpc" else "PID Water")
            sheets[sheet] = pd.read_csv(f)
        write_excel(ROOT / "excel" / b / f"{b}.xlsx", sheets)

    general = {}
    for name, f in (("Official Summary", "official_summary.csv"), ("Official Per Scenario", "official_per_scenario.csv"),
                    ("Pre-Conditioning Check", "preconditioning_check.csv"),
                    ("Loop Benchmark", "loop_benchmark_6_scenarios_2_days.csv"),
                    ("World Model Accuracy", "world_model_accuracy.csv")):
        if (ROOT / "csv" / f).exists():
            general[name] = pd.read_csv(ROOT / "csv" / f)
    write_excel(ROOT / "excel" / "summary.xlsx", general)
    (ROOT / "csv" / "README.txt").write_text(
        "General tables at this level; building_<n> folders hold per-building tables and the "
        f"{args.days}-day trajectories of both controllers. Policy: {args.weights}.\n")
    print("results ready in", ROOT.resolve())


if __name__ == "__main__":
    main()
